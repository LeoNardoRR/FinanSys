import csv
import io
import shutil
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, RedirectResponse, StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.database import engine, get_session
from app.models import Bank, Card, Category, Person, Transaction
from app.settings import DATA_DIR, DATABASE_PATH

router = APIRouter(prefix="/reports", tags=["reports"])


def _transactions(session: Session):
    return session.scalars(select(Transaction).order_by(Transaction.occurred_on.desc(), Transaction.id.desc())).all()


def _lookups(session: Session):
    def names(model):
        return {item.id: item.name for item in session.scalars(select(model)).all()}
    return names(Category), names(Bank), names(Card), names(Person)


@router.get("/transactions.csv", name="export_transactions_csv")
def export_transactions_csv(session: Session = Depends(get_session)):
    categories, banks, cards, people = _lookups(session)
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow(["Data", "Tipo", "Descrição", "Valor", "Categoria", "Conta", "Cartão", "Pessoa", "Parcela", "Observação"])
    for item in _transactions(session):
        writer.writerow([item.occurred_on.strftime("%d/%m/%Y"), item.kind, item.description,
                         str(item.amount).replace(".", ","), categories.get(item.category_id, ""),
                         banks.get(item.bank_id, ""), cards.get(item.card_id, ""), people.get(item.person_id, "Família"),
                         f"{item.installment_number or ''}/{item.installments_total or ''}", item.notes or ""])
    return StreamingResponse(iter([output.getvalue().encode("utf-8-sig")]), media_type="text/csv",
                             headers={"Content-Disposition": "attachment; filename=lancamentos_finansys.csv"})


@router.get("/transactions.xlsx", name="export_transactions_xlsx")
def export_transactions_xlsx(session: Session = Depends(get_session)):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise HTTPException(503, "Instale openpyxl para exportar Excel.") from exc
    categories, banks, cards, people = _lookups(session)
    book = Workbook()
    sheet = book.active
    sheet.title = "Lançamentos"
    headers = ["Data", "Tipo", "Descrição", "Valor", "Categoria", "Conta", "Cartão", "Pessoa", "Parcela", "Observação"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
    for item in _transactions(session):
        sheet.append([item.occurred_on, "Receita" if item.kind == "income" else "Despesa", item.description,
                      float(item.amount), categories.get(item.category_id, ""), banks.get(item.bank_id, ""),
                      cards.get(item.card_id, ""), people.get(item.person_id, "Família"),
                      f"{item.installment_number or ''}/{item.installments_total or ''}", item.notes or ""])
    sheet.column_dimensions["A"].width = 13
    sheet.column_dimensions["C"].width = 32
    sheet.column_dimensions["J"].width = 40
    for cell in sheet["D"][1:]:
        cell.number_format = 'R$ #,##0.00'
    stream = io.BytesIO()
    book.save(stream)
    stream.seek(0)
    return StreamingResponse(stream, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": "attachment; filename=lancamentos_finansys.xlsx"})


@router.get("/summary.pdf", name="export_summary_pdf")
def export_summary_pdf(session: Session = Depends(get_session)):
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
    except ImportError as exc:
        raise HTTPException(503, "Instale reportlab para exportar PDF.") from exc
    items = _transactions(session)
    income = sum((Decimal(i.amount) for i in items if i.kind == "income"), Decimal(0))
    expense = sum((Decimal(i.amount) for i in items if i.kind == "expense"), Decimal(0))
    stream = io.BytesIO()
    pdf = canvas.Canvas(stream, pagesize=A4)
    width, height = A4
    pdf.setTitle("Resumo financeiro FinanSys")
    pdf.setFont("Helvetica-Bold", 20)
    pdf.drawString(42, height - 55, "FinanSys - Resumo financeiro")
    pdf.setFont("Helvetica", 11)
    pdf.drawString(42, height - 82, f"Gerado em {datetime.now():%d/%m/%Y %H:%M}")
    pdf.drawString(42, height - 120, f"Receitas: R$ {income:,.2f}")
    pdf.drawString(42, height - 140, f"Despesas: R$ {expense:,.2f}")
    pdf.drawString(42, height - 160, f"Saldo: R$ {income - expense:,.2f}")
    y = height - 205
    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(42, y, "Últimos lançamentos")
    pdf.setFont("Helvetica", 9)
    for item in items[:28]:
        y -= 18
        if y < 50:
            pdf.showPage(); y = height - 50; pdf.setFont("Helvetica", 9)
        signal = "+" if item.kind == "income" else "-"
        pdf.drawString(42, y, f"{item.occurred_on:%d/%m/%Y}  {item.description[:48]}")
        pdf.drawRightString(width - 42, y, f"{signal} R$ {Decimal(item.amount):,.2f}")
    pdf.save()
    stream.seek(0)
    return StreamingResponse(stream, media_type="application/pdf",
                             headers={"Content-Disposition": "attachment; filename=resumo_finansys.pdf"})


@router.post("/backup", name="create_backup")
def create_backup():
    backup_dir = DATA_DIR / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    name = f"finansys_backup_{datetime.now():%Y%m%d_%H%M%S}.zip"
    path = backup_dir / name
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        if DATABASE_PATH.exists():
            archive.write(DATABASE_PATH, arcname="finansys.db")
    return RedirectResponse(f"/settings?backup={name}", 303)


@router.get("/backup/{filename}", name="download_backup")
def download_backup(filename: str):
    safe_name = Path(filename).name
    path = DATA_DIR / "backups" / safe_name
    if not path.is_file() or path.suffix.lower() != ".zip":
        raise HTTPException(404, "Backup não encontrado.")
    return FileResponse(path, filename=safe_name, media_type="application/zip")


@router.post("/restore", name="restore_backup")
async def restore_backup(file: UploadFile = File(...)):
    if not file.filename or Path(file.filename).suffix.lower() != ".zip":
        raise HTTPException(422, "Selecione um backup .zip criado pelo FinanSys.")
    content = await file.read()
    with NamedTemporaryFile(suffix=".zip", delete=False) as temp:
        temp.write(content)
        temp_path = Path(temp.name)
    try:
        with ZipFile(temp_path) as archive:
            if "finansys.db" not in archive.namelist():
                raise HTTPException(422, "O arquivo não contém um banco FinanSys.")
            backup_dir = DATA_DIR / "backups"
            backup_dir.mkdir(parents=True, exist_ok=True)
            if DATABASE_PATH.exists():
                shutil.copy2(DATABASE_PATH, backup_dir / f"antes_restauracao_{datetime.now():%Y%m%d_%H%M%S}.db")
            engine.dispose()
            with archive.open("finansys.db") as source, open(DATABASE_PATH, "wb") as target:
                shutil.copyfileobj(source, target)
    finally:
        temp_path.unlink(missing_ok=True)
    return RedirectResponse("/settings?restored=1", 303)
